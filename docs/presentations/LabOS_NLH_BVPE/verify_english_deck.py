"""Validate the English deliverables and export reusable vector figures."""
from __future__ import annotations

import hashlib
import html
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image,ImageDraw
from pptx import Presentation

import build_presentation as g

ROOT=Path(__file__).resolve().parent/"English"


def clean(text):return re.sub(r"\s+","",text)


def render_previews(pdf_path):
    with tempfile.TemporaryDirectory(prefix="labos-slides-") as directory:
        prefix=Path(directory)/"slide"
        subprocess.run(["pdftoppm","-scale-to","1440","-png",str(pdf_path),str(prefix)],check=True,capture_output=True)
        return [Image.open(p).convert("RGB") for p in sorted(Path(directory).glob("slide-*.png"))]


def export_svg(scene,path,view):
    x,y,w,h=view
    parts=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{x} {y} {w} {h}" width="{w}" height="{h}">',
           '<title>'+html.escape(scene["title"])+"</title>"]
    for o in scene["ops"]:
        k=o["kind"]
        fill="#"+o["fill"] if o.get("fill") else "none"
        stroke="#"+o["border"] if o.get("border") else "none"
        sw=o.get("lw",1)
        dash=' stroke-dasharray="5 4"' if o.get("dash") else ""
        if k=="rect":
            parts.append(f'<rect x="{o["x"]}" y="{o["y"]}" width="{o["w"]}" height="{o["h"]}" rx="{o["radius"]}" fill="{fill}" stroke="{stroke}" stroke-width="{sw}"{dash}/>')
        elif k=="circle":
            parts.append(f'<circle cx="{o["x"]}" cy="{o["y"]}" r="{o["r"]}" fill="{fill}" stroke="{stroke}" stroke-width="{sw}"/>')
        elif k=="poly":
            pts=" ".join(f"{xx},{yy}" for xx,yy in o["points"])
            parts.append(f'<polygon points="{pts}" fill="{fill}" stroke="{stroke}" stroke-width="{sw}"/>')
        elif k=="line":
            parts.append(f'<line x1="{o["x1"]}" y1="{o["y1"]}" x2="{o["x2"]}" y2="{o["y2"]}" stroke="#{o["color"]}" stroke-width="{sw}"{dash}/>')
        elif k=="text":
            for j,line in enumerate(o["lines"]):
                tw=g.width(line,o["size"],o["bold"])
                xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
                yy=o["y"]+o["size"]*.93+j*o["size"]*o["leading"]
                weight=700 if o["bold"] else 400
                parts.append(f'<text x="{xx}" y="{yy}" font-family="DejaVu Sans,Arial,sans-serif" font-size="{o["size"]}" font-weight="{weight}" fill="#{o["color"]}">{html.escape(line)}</text>')
    parts.append("</svg>")
    path.write_text("\n".join(parts),encoding="utf-8")


def main():
    scenes=json.loads((ROOT/"slides_scene_EN.json").read_text())
    pdf=pymupdf.open(ROOT/"LabOS_Technical_Overview_EN.pdf")
    prs=Presentation(ROOT/"LabOS_Technical_Overview_EN.pptx")
    assert len(scenes)==len(pdf)==len(prs.slides)==22
    previews=render_previews(ROOT/"LabOS_Technical_Overview_EN.pdf")
    assert len(previews)==22
    metadata=json.dumps(pdf.metadata)+str(prs.core_properties.title)+str(prs.core_properties.subject)
    assert not re.search(r"[\u3400-\u9fff]",metadata)
    errors=[]; overlap=[]; thumbs=[]; tb=0; vectors=0
    for i,(scene,page,slide) in enumerate(zip(scenes,pdf,prs.slides),1):
        ppt="\n".join(s.text for s in slide.shapes if s.has_text_frame)
        note=slide.notes_slide.notes_text_frame.text
        text=page.get_text()
        assert not re.search(r"[\u3400-\u9fff]",ppt+note+text),f"non-English text on page {i}"
        assert len(note)>150
        bounds=[]
        for oi,o in enumerate(scene["ops"]):
            if o["kind"]!="text":continue
            for li,line in enumerate(o["lines"]):
                if clean(line) and clean(line) not in clean(text):errors.append(f"PDF page {i}: {line}")
                if clean(line) and clean(line) not in clean(ppt):errors.append(f"PPT slide {i}: {line}")
                tw=g.width(line,o["size"],o["bold"])
                xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
                yy=o["y"]+li*o["size"]*o["leading"]
                if line:bounds.append((oi,line,(xx,yy,xx+tw,yy+o["size"])))
        for bi,a in enumerate(bounds):
            for b in bounds[bi+1:]:
                if a[0]==b[0]:continue
                ax,ay,ar,ab=a[2];bx,by,br,bb=b[2]
                ix=min(ar,br)-max(ax,bx);iy=min(ab,bb)-max(ay,by)
                if ix>10 and iy>6:overlap.append(dict(slide=i,a=a[1],b=b[1],intersection=[round(ix,1),round(iy,1)]))
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text:tb+=1
            else:vectors+=1
            assert int(sh.shape_type)!=13
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines",[]):
                for sp in line["spans"]:
                    x0,y0,x1,y1=sp["bbox"]
                    if min(x0,y0)<-1 or x1>page.rect.width+1 or y1>page.rect.height+1:errors.append(f"page {i} out of bounds: {sp['text']}")
        im=previews[i-1]
        if i in [3,6,8,10,13,17,18]:im.save(ROOT/f"Figure_preview_{i:02d}.png")
        im.thumbnail((480,270));tile=Image.new("RGB",(500,304),"#E2E8F0")
        tile.paste(im,((500-im.width)//2,8));ImageDraw.Draw(tile).text((13,280),f"{i:02d}",fill="#11233E");thumbs.append(tile)
    if errors:raise AssertionError("\n".join(errors))
    sheet=Image.new("RGB",(1500,304*((len(thumbs)+2)//3)),"#D2DDE9")
    for i,t in enumerate(thumbs):sheet.paste(t,((i%3)*500,(i//3)*304))
    sheet.save(ROOT/"Slide_Overview_EN.png")
    figures=ROOT/"Figures";figures.mkdir(exist_ok=True)
    definitions=[(3,"01_System_Architecture",(45,177,1190,500)),
                 (6,"02_Experiment_Topology",(45,178,1190,438)),
                 (8,"03_Timing_and_References",(45,183,1190,418)),
                 (10,"04_Nonlinear_Hall_Circuit",(45,185,1190,433)),
                 (13,"05_BVPE_Optical_and_Readout",(45,185,1190,431)),
                 (18,"06_Execution_and_Quality",(45,190,1190,418))]
    for n,name,view in definitions:
        export_svg(scenes[n-1],figures/(name+".svg"),view)
        x,y,w,h=view;clip=pymupdf.Rect(x*.75,y*.75,(x+w)*.75,(y+h)*.75)
        out=pymupdf.open();p=out.new_page(width=clip.width,height=clip.height)
        p.show_pdf_page(p.rect,pdf,n-1,clip=clip)
        out.save(figures/(name+".pdf"));out.close()
    report=dict(status="PASS" if not overlap else "REVIEW_OVERLAPS",pages=22,language="English throughout slides, speaker notes and metadata",
                native_text_boxes=tb,native_vector_shapes_and_lines=vectors,slide_raster_images=0,
                missing_text=0,out_of_page_text=0,text_overlap_candidates=overlap,
                speaker_notes="22/22",standalone_vector_figures=len(definitions),
                reference_links=sum(len(p.get_links()) for p in pdf))
    (ROOT/"QA_Report_EN.json").write_text(json.dumps(report,indent=2))
    files=[ROOT/name for name in ["LabOS_Technical_Overview_EN.pptx","LabOS_Technical_Overview_EN.pdf","Speaker_Notes_EN.md","Slide_Overview_EN.png","QA_Report_EN.json","README.md","slides_scene_EN.json"] if (ROOT/name).exists()]
    files.extend(sorted(ROOT.glob("Figure_preview_*.png")))
    files.extend(sorted(figures.glob("*")))
    (ROOT/"SHA256SUMS").write_text("".join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(ROOT).as_posix()}\n" for p in files))
    bundle=ROOT/"LabOS_Technical_Overview_EN_Bundle.zip"
    with zipfile.ZipFile(bundle,"w",compression=zipfile.ZIP_DEFLATED) as archive:
        for p in files+[ROOT/"SHA256SUMS"]:archive.write(p,p.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(bundle) as archive:
        assert archive.testzip() is None
        for p in files:
            assert archive.read(p.relative_to(ROOT).as_posix())==p.read_bytes()
    print(json.dumps(report,indent=2))
    print(f"Bundle integrity: PASS ({len(files)+1} files)")


if __name__=="__main__":main()
