"""Export a native-vector transport workflow, its preview and delivery bundle."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path

import pymupdf
from pptx import Presentation

import build_presentation as g
from build_system_architecture import monitor, panel, text
from verify_english_deck import export_svg

ROOT=Path(__file__).resolve().parent/"Transport_Closed_Loop"
ROOT.mkdir(exist_ok=True)
W,H=3000,2000
C=g.C
NAME="LabOS_Transport_Closed_Loop_EN"


class Scene(g.Slide):
    def __init__(self):
        self.title="LabOS Transport: Calibration, Feedback and Physics-Guided Planning"
        self.section="Transport experiment workflow"
        self.note=("This is a transport workflow design, complementing the general LabOS architecture. "
                   "The left column separates the server-to-worker HTTP(S) link, local driver APIs and physical instrument interfaces. "
                   "A master lock-in supplies sine excitation through a series resistor and reads Vxx at the fundamental; "
                   "a follower reads Vxy at a selected harmonic using a physical reference. Dedicated calibration stages reuse "
                   "the configured readout paths; no extra simultaneous measurement channel or automatic rewiring is assumed. "
                   "Autonomous calibration adjusts bounded software settings, repeats references and issues a versioned calibration record. "
                   "Raw X/Y and reference phase are retained, rather than auto-phasing an unknown second-harmonic signal to zero. "
                   "The central loop qualifies each completed point or chunk, proposes a revision and validates a new stage version before dispatch. "
                   "Recoverable quality problems can return to calibration; faults or uncertain commands require pause and reconciliation. "
                   "The server-local theory/compute coupling is explicitly planned. It receives versioned qualified-data snapshots, "
                   "compares transport, thermal and circuit hypotheses and ranks informative candidate measurements within the operating envelope. "
                   "This is not a claim of deployed transport autonomy or measured material results.")
        self.source="Original LabOS workflow design; supporting primary sources are linked in the bilingual companion document."
        self.ops=[];self.errors=[];self.dark=False
        self.rect(0,0,W,H,C["white"])


def node(s,x,y,w,h,title,body,color,fill=None,ts=28,bs=22):
    s.rect(x,y,w,h,fill or C["white"],color,radius=12,lw=1.7)
    text(s,x+22,y+16,w-44,45,title,ts,color,True)
    if body:text(s,x+22,y+66,w-44,h-78,body,bs)


def rack(s,x,y,w,title):
    s.rect(x,y,w,48,C["pale"],C["blue"],radius=5,lw=1.2)
    s.rect(x+10,y+8,w-77,32,C["navy"],radius=3)
    text(s,x+20,y+12,w-99,29,title,21,C["white"],True)
    s.circle(x+w-33,y+24,10,C["white"],C["blue"],1.2)
    s.circle(x+w-13,y+36,3,C["green"])


def build():
    s=Scene()
    text(s,60,35,2880,73,"LabOS | A Transport Experiment That Calibrates and Learns",47,C["navy"],True)
    text(s,63,122,2870,41,"Device communication → autonomous calibration → live feedback → physics-guided next experiment",27,C["muted"])
    for xx,ww,t in [(60,820,"A / DEVICE COMMUNICATION"),(1010,980,"B / MEASUREMENT AND FEEDBACK"),(2120,820,"C / SCIENTIFIC NEXT-STEP DESIGN")]:
        text(s,xx,224,ww,39,t,27,C["navy"],True)

    # Communication layers: MCP/API, outbound worker channel, driver API, wire protocol.
    panel(s,60,300,820,1360,"Commands down • evidence back",C["blue"],C["bg"],29)
    node(s,90,375,750,115,"Central tools → bounded command","command_id / stage_id / device_id / deadline",C["purple"],C["plum"],26,21)
    node(s,90,550,750,82,"PC-initiated HTTP(S)","",C["purple"],C["white"],27)
    s.arrow([(440,490),(440,550)],C["purple"],2.4,head=9)
    s.arrow([(440,632),(440,725)],C["purple"],2.4,head=9)
    text(s,112,650,300,61,"Poll + heartbeat\nResults + raw artifacts",21,C["purple"])
    s.rect(90,725,750,160,C["ice"],C["teal"],radius=11,lw=1.7)
    monitor(s,112,745,65,45,C["teal"])
    text(s,200,744,610,42,"PC i … PC N  /  local PC Agent",29,C["teal"],True)
    text(s,113,803,699,64,"Device sessions + command journal + watchdog\nReadbacks / flags / timestamps / local raw cache",22)
    s.arrow([(840,816),(858,816),(858,433),(840,433)],C["green"],2.1,head=8)
    s.arrow([(440,885),(440,955)],C["teal"],2.2,head=8)
    node(s,90,955,750,112,"Driver adapters / reviewed local code","VISA I/O or vendor SDK; SCPI only where supported",C["teal"],C["white"],27,20)
    s.arrow([(415,1067),(415,1150)],C["blue"],2.3,head=9)
    s.arrow([(445,1150),(445,1067)],C["blue"],2.3,head=9)
    text(s,470,1090,360,32,"GPIB / USB / serial / TCP-IP",21,C["blue"])
    rack(s,90,1150,750,"Lock-in master: sine OUT + Vxx(1f)")
    rack(s,90,1209,750,"Lock-in follower: selected Vxy(2f)")
    rack(s,90,1268,750,"attoDRY: T / B    |    2400: Vg / Ig")
    text(s,92,1365,747,38,"Physical signal and reference cables",26,C["blue"],True)
    s.rect(100,1430,158,60,C["plum"],C["purple"],radius=7)
    text(s,110,1447,138,30,"Sine OUT",23,C["purple"],True,align="center")
    s.line(258,1460,300,1460,C["purple"],2)
    pts=[(300,1460),(311,1460)]+[(316+i*9,1460+(-8 if i%2==0 else 8)) for i in range(8)]+[(391,1460),(405,1460)]
    for a,b in zip(pts,pts[1:]):s.line(*a,*b,C["purple"],2)
    text(s,307,1417,97,30,"Rs",22,C["purple"],True,align="center")
    s.arrow([(405,1460),(461,1460)],C["purple"],2.1,head=8)
    s.rect(461,1430,235,60,C["pale"],C["blue"],radius=7)
    text(s,472,1446,213,31,"Sample + return",23,C["blue"],True,align="center")
    s.arrow([(126,1550),(789,1550)],C["amber"],2.4,True,head=9)
    text(s,108,1507,690,33,"Master f_I → follower REF; select harmonic",22,C["amber"],True)
    text(s,108,1580,690,54,"Network timing and physical phase reference\nare separate parts of the experiment.",21,C["muted"])

    # The central point/chunk loop; each revision becomes a new validated stage.
    node(s,1010,300,980,130,"1  Define the transport campaign","Sample / wiring / T, B, Vg / current and frequency envelope",C["purple"],C["plum"],29,22)
    node(s,1010,500,980,180,"2  Autonomous calibration","Verify actual current, reference and phase\nChoose range / filtering; repeat baseline checks\nReuse valid calibration or issue a new calibration_id",C["amber"],C["sand"],30,23)
    node(s,1010,760,980,170,"3  Prepare + acquire the next point","Join actual T/B/Vg readiness; settle after setting changes\nRecheck calibration validity at the observed state\nSave Vxx(1f), Vxy(2f), X/Y, readbacks and timestamps",C["teal"],C["ice"],29,21)
    node(s,1010,990,980,150,"4  Live qualification","Reference lock / overload / freshness / drift / leakage\nRecoverable QC issue → recalibrate; hard fault → pause",C["green"],C["mint"],29,23)
    node(s,1010,1230,980,150,"5  Central feedback + scientific analysis","Fit amplitude / phase / drive dependence; compare controls\nPropose a repeat, a new setting or a discriminating test",C["purple"],C["plum"],29,23)
    node(s,1010,1470,980,130,"6  Validate update / close the campaign","Limits + locks + versioned update / recorded closure",C["teal"],C["ice"],29,23)
    for yy,ny,col in [(430,500,C["purple"]),(680,760,C["teal"]),(930,990,C["teal"]),(1140,1230,C["green"]),(1380,1470,C["purple"])]:
        s.arrow([(1500,yy),(1500,ny)],col,2.8,head=10)
    text(s,1525,1170,430,33,"Qualified evidence",21,C["green"],True)
    s.arrow([(1010,1535),(940,1535),(940,590),(1010,590)],C["purple"],2.7,head=10)
    s.arrow([(1990,1040),(2040,1040),(2040,590),(1990,590)],C["amber"],2.5,head=10)
    s.arrow([(1500,1600),(1500,1640)],C["green"],2.4,head=9)
    s.rect(1010,1640,980,56,C["mint"],C["green"],radius=8,lw=1.4)
    text(s,1032,1654,936,32,"Objective met / budget reached → archive + close",24,C["green"],True,align="center")

    # Explicitly planned computational physics service, separate from the local LLM.
    s.rect(2120,300,820,660,C["sand"],C["amber"],radius=14,lw=1.8,dash=True)
    text(s,2142,321,776,43,"PLANNED • SERVER-LOCAL THEORY",28,C["amber"],True)
    text(s,2142,374,776,32,"Theory MCP / numerical jobs / versioned models",22,C["amber"])
    theory=[(435,"Receive a qualified-data snapshot","X/Y, I, f, T/B/Vg + calibration + snapshot IDs"),
            (559,"Run models + estimate uncertainty","Transport / circuit / thermal-response hypotheses"),
            (683,"Compare predictions with experiment","Residuals, validity range and competing explanations"),
            (807,"Rank informative next measurements","Model discrimination / uncertainty / time / constraints")]
    for yy,t,b in theory:
        node(s,2143,yy,774,102,t,b,C["amber"],C["white"],25,20)
        if yy!=807:s.arrow([(2530,yy+102),(2530,yy+124)],C["amber"],1.7,head=7)
    s.arrow([(1990,1278),(2090,1278),(2090,461),(2143,461)],C["purple"],2.1,True,head=9)
    s.arrow([(2143,866),(2065,866),(2065,1335),(1990,1335)],C["green"],2.1,True,head=9)

    panel(s,2120,1040,820,560,"What changes after feedback?",C["blue"],C["pale"],29)
    rows=[(1113,"Noisy but stable","Longer averaging / suitable filtering",C["green"]),
          (1214,"Overload / reference drift","Correct range / reference; reacquire",C["amber"]),
          (1315,"Nonlinearity / possible heating","Lower drive; add frequency / return controls",C["purple"]),
          (1416,"Fault / uncertain command","Pause, local closure and state reconciliation",C["red"])]
    for yy,t,b,col in rows:
        s.rect(2143,yy,774,87,C["white"],C["line"],radius=7)
        text(s,2159,yy+9,741,31,t,24,col,True)
        text(s,2159,yy+48,741,30,b,22)
    text(s,2143,1530,772,53,"Every revision preserves raw evidence and passes\nthrough the same execution authority.",21,C["blue"],True)

    for x,title,body,col,fill in [(60,"LOCAL DEVICE LOOPS","Reference / readout / local fault response\nNo LLM in the timing-critical path",C["teal"],C["ice"]),
                                 (1040,"CENTRAL POINT / CHUNK FEEDBACK","Event-driven QC and bounded stage revisions\nDashboard follows state and qualified data",C["purple"],C["plum"]),
                                 (2020,"PLANNED ASYNCHRONOUS THEORY","Compute budget + uncertainty + snapshot checks\nA local LLM is not a numerical physics solver",C["amber"],C["sand"])]:
        node(s,x,1740,920,143,title,body,col,fill,25,22)
    text(s,60,1920,2880,35,"Transport workflow design; the theory/compute coupling is planned. No measured curves or material mechanism claims are implied.",23,C["muted"])
    return s


def validate(scene,pptx,pdf):
    prs=Presentation(pptx);doc=pymupdf.open(pdf)
    assert len(prs.slides)==len(doc)==1
    ppt="\n".join(sh.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    pdftext=doc[0].get_text()
    errors=list(scene.errors);bounds=[]
    norm=lambda s:re.sub(r"\s+","",s)
    for oi,o in enumerate(scene.ops):
        if o["kind"]!="text":continue
        for j,line in enumerate(o["lines"]):
            for kind,content in [("PPTX",ppt),("PDF",pdftext)]:
                if norm(line) not in norm(content):errors.append(f"{kind} missing: {line}")
            tw=g.width(line,o["size"],o["bold"])
            xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
            yy=o["y"]+j*o["size"]*o["leading"]
            if line:bounds.append((oi,line,(xx,yy,xx+tw,yy+o["size"])))
    for i,a in enumerate(bounds):
        for b in bounds[i+1:]:
            if a[0]==b[0]:continue
            ax,ay,ar,ab=a[2];bx,by,br,bb=b[2]
            if min(ar,br)-max(ax,bx)>8 and min(ab,bb)-max(ay,by)>5:errors.append(f"Text overlap: {a[1]} / {b[1]}")
    assert not re.search(r"[\u3400-\u9fff]",ppt+prs.slides[0].notes_slide.notes_text_frame.text)
    assert not any(int(sh.shape_type)==13 for sh in prs.slides[0].shapes)
    for block in doc[0].get_text("dict")["blocks"]:
        for line in block.get("lines",[]):
            for span in line["spans"]:
                x0,y0,x1,y1=span["bbox"]
                if min(x0,y0)<-1 or x1>W*.5+1 or y1>H*.5+1:errors.append(f"PDF span outside page: {span['text']}")
    for required in ["Autonomous calibration","PC-initiated HTTP(S)","PLANNED • SERVER-LOCAL THEORY","Validate update","calibration_id","budget reached"]:assert required in ppt
    if errors:raise AssertionError("\n".join(errors))
    return dict(status="PASS",pages=1,language="English",native_objects=len(prs.slides[0].shapes),raster_images_in_pptx=0,
                missing_text=0,text_overlaps=0,out_of_page_text=0,layout_bounds="PASS",theory_coupling="explicitly planned",device_control_code_changes=False)


def main():
    g.W=W;g.H=H;g.SCALE=.5;g.ROOT=ROOT;g.SLIDES.clear()
    s=build();g.SLIDES.append(s)
    if s.errors:raise AssertionError("\n".join(s.errors))
    pptx=g.render_pptx(NAME+".pptx",s.title,"Communication, autonomous calibration, feedback and planned physics guidance",notes_source_label="Reference: ")
    pdf=g.render_pdf(NAME+".pdf",s.title)
    record=dict(title=s.title,section=s.section,note=s.note,source=s.source,ops=s.ops)
    (ROOT/"workflow_scene.json").write_text(json.dumps(record,indent=2,ensure_ascii=False)+"\n")
    export_svg(record,ROOT/(NAME+".svg"),(0,0,W,H))
    report=validate(s,pptx,pdf)
    for size,name in [(6000,NAME),(2100,"Workflow_Preview")]:
        subprocess.run(["pdftoppm","-singlefile","-scale-to",str(size),"-png",str(pdf),str(ROOT/name)],capture_output=True,check=True)
    (ROOT/"QA_Report.json").write_text(json.dumps(report,indent=2)+"\n")
    files=sorted(p for p in ROOT.iterdir() if p.is_file() and p.name not in {"SHA256SUMS",NAME+"_Bundle.zip"})
    (ROOT/"SHA256SUMS").write_text("".join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}\n" for p in files))
    with zipfile.ZipFile(ROOT/(NAME+"_Bundle.zip"),"w",compression=zipfile.ZIP_DEFLATED) as z:
        for p in files+[ROOT/"SHA256SUMS"]:z.write(p,p.name)
    with zipfile.ZipFile(ROOT/(NAME+"_Bundle.zip")) as z:
        assert z.testzip() is None
        for p in files:assert z.read(p.name)==p.read_bytes()
    print(json.dumps(report,indent=2));print(f"Bundle integrity: PASS ({len(files)+1} files)")


if __name__=="__main__":main()
