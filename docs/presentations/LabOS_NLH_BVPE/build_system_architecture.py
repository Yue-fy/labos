"""Build a standalone LabOS architecture figure using editable native vectors."""
from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import zipfile
from pathlib import Path

import pymupdf
from pptx import Presentation

import build_presentation as g
from verify_english_deck import export_svg

ROOT=Path(__file__).resolve().parent/"System_Architecture"
ROOT.mkdir(exist_ok=True)
W,H=2560,1920
C=g.C


class Scene(g.Slide):
    def __init__(self):
        self.title="LabOS System Architecture"
        self.section="Skill-driven scientific runtime and distributed execution"
        self.note=("This redesigned LabOS architecture uses three interchangeable PC symbols for an arbitrary fleet, "
                   "with no fixed assignment of scientific functions to PCs. One central LabOS server hosts "
                   "scientific AI reasoning, specialist sub-agents, knowledge and memory, a skill library, "
                   "MCP host/client components and tool connectors, a capability registry, "
                   "and deterministic coordination. A server-side model router prefers a remote cloud model and falls back "
                   "to the server's local LLM when cloud connectivity, the API or its quota is unavailable. "
                   "The agents preserve experiment context and tool permissions across model changes. "
                   "Each instrument-host PC runs one persistent execution worker, called a PC Agent. "
                   "The worker does not require a local large language model. It owns instrument sessions, bounded adapters, "
                   "a command journal, watchdog and artifact outbox. An operator computer can use only a browser. "
                   "All laboratory PCs initiate their HTTP(S) connections; command arrows represent payload flow, "
                   "not inbound network sessions. Physical device interfaces remain local to the instrument host. "
                   "Skills package reusable scientific procedures and supporting references or scripts. MCP exposes tools, "
                   "resources and prompts; the LabOS control connector routes instrument actions through the deterministic coordinator. "
                   "Hardware reference and trigger cables are separate from this control network. "
                   "This figure describes the scalable LabOS system design, rather than an inventory of currently commissioned drivers.")
        self.source="Original LabOS architecture; protocol references: https://agentskills.io/home and https://modelcontextprotocol.io/docs/learn/architecture"
        self.ops=[];self.errors=[];self.dark=False
        self.rect(0,0,W,H,C["white"])


def text(s,x,y,w,h,t,size=22,color=None,bold=False,align="left"):
    s.text(x,y,w,h,t,size,color or C["ink"],bold,align,leading=1.25)


def panel(s,x,y,w,h,title,color,fill,title_size=28):
    s.rect(x,y,w,h,fill,color,radius=15,lw=1.8)
    text(s,x+22,y+18,w-44,45,title,title_size,color,True)


def monitor(s,x,y,w=65,h=48,color=None):
    color=color or C["blue"]
    s.rect(x,y,w,h,C["pale"],color,radius=4,lw=2)
    s.rect(x+7,y+7,w-14,h-14,C["white"])
    s.line(x+w/2,y+h,x+w/2,y+h+10,color,2)
    s.line(x+w*.21,y+h+12,x+w*.79,y+h+12,color,3)


def server_icon(s,x,y):
    s.rect(x,y,63,75,C["pale"],C["blue"],radius=7,lw=2)
    for k in range(3):
        yy=y+10+k*21
        s.rect(x+8,yy,47,15,C["white"],C["blue"],radius=3,lw=1)
        s.circle(x+46,yy+7.5,2.5,C["green"])
        s.line(x+14,yy+7.5,x+33,yy+7.5,C["blue"],1.5)


def database(s,x,y,w=72):
    for offset in (35,17,0):
        s.rect(x,y+offset,w,29,C["white"],C["green"],radius=10,lw=1.7)
    pts=[(x+w/2+w/2*math.cos(k*math.pi/12),y+8+8*math.sin(k*math.pi/12)) for k in range(24)]
    s.poly(pts,C["mint"],C["green"],1.5)


def card(s,x,y,w,h,title,body,color=C["teal"],fill=C["white"],title_size=22,body_size=18):
    s.rect(x,y,w,h,fill,color,radius=8,lw=1.3)
    text(s,x+12,y+10,w-24,31,title,title_size,color,True)
    if body:text(s,x+12,y+45,w-24,h-50,body,body_size)


def instrument(s,x,y,w,title,sub="",h=52):
    s.rect(x,y,w,h,C["pale"],C["blue"],radius=5,lw=1.3)
    s.rect(x+9,y+8,w-65,h-16,C["navy"],radius=3)
    size=17
    while g.width(title,size,True)>w-83:size-=.5
    text(s,x+17,y+13,w-83,26,title,size,C["white"],True)
    s.circle(x+w-30,y+h/2,10,C["white"],C["blue"],1.3)
    s.circle(x+w-11,y+h-10,2.8,C["green"])
    if sub:text(s,x+12,y+h+5,w-24,23,sub,15,C["muted"])


def build():
    s=Scene()
    text(s,60,31,2430,65,"LabOS | Scientific Skills, Connected Instruments",45,C["navy"],True)
    text(s,62,99,2410,34,"Agents decide • skills guide workflows • MCP connects tools • local workers execute",24,C["muted"])

    # Operator and remote inference service sit outside the one-server boundary.
    s.rect(60,151,1010,61,C["sand"],C["amber"],radius=12,lw=1.6)
    monitor(s,80,161,48,31,C["amber"])
    text(s,151,164,894,34,"Operator workspace: objective, setup, review and authorization",24,C["amber"],True)
    s.arrow([(565,212),(565,240)],C["amber"],2.4,head=9)
    s.rect(2100,151,400,61,C["plum"],C["purple"],radius=12,lw=1.5,dash=True)
    text(s,2118,168,364,33,"Cloud LLM API • remote",24,C["purple"],True,align="center")
    s.arrow([(2284,240),(2284,212)],C["purple"],2.2,head=8)
    s.arrow([(2314,212),(2314,240)],C["green"],2.2,head=8)

    s.rect(60,240,2440,1040,C["bg"],C["blue"],radius=20,lw=2.6)
    server_icon(s,85,259)
    text(s,171,262,1350,47,"ONE CENTRAL SERVER  /  LABOS RUNTIME",34,C["navy"],True)
    text(s,172,307,2240,30,"Scientific reasoning, reusable methods, tool access, model routing and persistent execution state",22,C["muted"])

    # Four capability services surround a scientific agent kernel and its tool plane.
    panel(s,100,370,560,280,"Knowledge & memory",C["green"],C["mint"],29)
    database(s,558,389,72)
    text(s,122,436,510,25,"SEMANTIC KNOWLEDGE",18,C["green"],True)
    text(s,122,471,510,57,"Literature / manuals / local code\nSetup definitions + calibrations",22)
    text(s,122,546,510,25,"EPISODIC EXPERIMENT MEMORY",18,C["green"],True)
    text(s,122,581,510,57,"Runs / artifacts / faults / decisions\nVersioned evidence + provenance",21)

    panel(s,800,370,1040,280,"Scientific agent runtime",C["purple"],C["plum"],32)
    s.rect(826,432,988,53,C["white"],C["purple"],radius=9,lw=1.6)
    text(s,842,446,956,34,"LabOS control agent / task supervisor",26,C["purple"],True,align="center")
    for i,t in enumerate(["Planning agent","Analysis agent","Diagnostics agent"]):
        xx=826+i*337
        s.rect(xx,513,314,45,C["white"],"C9B9EA",radius=7,lw=1.2)
        text(s,xx+10,526,294,29,t,21,C["purple"],True,align="center")
        s.arrow([(xx+157,485),(xx+157,513)],C["purple"],1.7,head=7)
    text(s,827,590,986,33,"Observe → frame → select skills → delegate → review → update",23,C["purple"],True,align="center")
    s.arrow([(660,493),(800,493)],C["green"],2.2,head=9)
    s.arrow([(800,519),(660,519)],C["green"],2.2,head=9)
    text(s,668,447,124,45,"Retrieve / persist",16,C["green"],True,align="center")

    panel(s,1900,370,560,280,"Model runtime / routing",C["purple"],C["plum"],28)
    for yy,t,sub in [(435,"Cloud API • preferred","Remote inference backend"),(517,"Local LLM • on this server","Fallback inference backend")]:
        s.rect(1922,yy,516,61,C["white"],C["purple"],radius=8,lw=1.3)
        text(s,1937,yy+6,486,31,t,23,C["purple"],True)
        text(s,1937,yy+37,486,24,sub,17,C["muted"])
    s.arrow([(2180,496),(2180,517)],C["purple"],2,head=7)
    text(s,1923,600,514,43,"Internet / API / quota failure → local model\nContext + tool permissions stay consistent",18,C["purple"])
    s.arrow([(1840,464),(1900,464)],C["purple"],2.2,head=8)
    s.arrow([(1900,489),(1840,489)],C["purple"],2.2,head=8)
    text(s,1843,424,54,26,"LLM",17,C["purple"],True,align="center")

    panel(s,100,755,560,275,"Skill library",C["amber"],C["sand"],30)
    skill_names=["Brainstorming","Experiment design","Calibration + QC","Instrument integration","Analysis + reporting","Troubleshooting"]
    for i,t in enumerate(skill_names):
        xx=122+(i%2)*258;yy=825+(i//2)*48
        s.rect(xx,yy,246,37,C["white"],"E3C995",radius=6,lw=1.1)
        text(s,xx+8,yy+8,230,25,t,17,C["amber"],True,align="center")
    text(s,122,981,511,41,"SKILL.md + references / scripts\nDiscover → load on demand → follow workflow",17,C["amber"])
    s.arrow([(660,861),(728,861),(728,579),(800,579)],C["amber"],2.2,head=9)
    text(s,664,698,58,48,"Load\nskill",16,C["amber"],True,align="center")

    # MCP is a software-tool boundary, distinct from instrument-host networking.
    panel(s,800,755,1040,275,"Tool access  /  MCP host + clients",C["blue"],C["pale"],28)
    for xx,t,body in [(826,"Knowledge MCP","Search / read / retrieve"),(1163,"Analysis MCP","Fit / compare / visualize"),(1500,"LabOS Control MCP","Request / status / cancel")]:
        s.rect(xx,832,314,102,C["white"],C["blue"],radius=8,lw=1.3)
        text(s,xx+12,845,290,31,t,23,C["blue"],True,align="center")
        text(s,xx+12,887,290,28,body,17,C["muted"],align="center")
    text(s,827,958,985,28,"Connector services expose tools, resources and prompts through a common interface.",20,C["blue"])
    text(s,827,991,985,27,"Additional connectors: papers, files, code, reporting and software integration",18,C["muted"])
    s.arrow([(1320,650),(1320,755)],C["blue"],2.6,head=10)
    text(s,1343,684,466,31,"Discover tools / invoke capability",21,C["blue"],True)
    s.arrow([(1280,755),(1280,650)],C["green"],2.6,head=10)
    text(s,845,684,410,31,"Structured results → evidence",20,C["green"],True,align="right")

    panel(s,1900,755,560,275,"Capability + setup registry",C["teal"],C["ice"],27)
    text(s,1923,824,514,141,"PC identity → device bindings\nOperations + parameter schemas\nUnits / limits + calibration\nShared resources + reservations\nRun / stage / point context",22)
    text(s,1923,985,514,29,"PC roles follow registered capabilities",20,C["teal"],True)
    s.arrow([(1840,868),(1900,868)],C["teal"],2,head=8)
    s.arrow([(1900,893),(1840,893)],C["teal"],2,head=8)

    # Every actuation request enters the same deterministic authority.
    s.arrow([(1657,1030),(1657,1110)],C["purple"],2.7,head=10)
    text(s,1680,1059,670,31,"LabOS control tools → validated execution",22,C["purple"],True)
    s.arrow([(1080,1110),(1080,1030)],C["green"],2.7,head=10)
    text(s,1110,1060,497,28,"Results / status / artifact handles",21,C["green"],True)
    panel(s,100,1110,2360,140,"Execution authority  /  deterministic stage coordinator",C["teal"],C["ice"],29)
    parts=["Validate units + limits","Reserve resources","Join READY states","Track command outcome","Qualify + archive data"]
    for i,t in enumerate(parts):
        xx=122+i*468
        s.rect(xx,1178,449,47,C["white"],C["teal"],radius=7,lw=1.2)
        text(s,xx+10,1192,429,29,t,21,C["teal"],True,align="center")

    # Representative, interchangeable PCs depict the entire fleet, rather than roles.
    text(s,63,1296,1050,28,"PC-INITIATED HTTP(S)  •  commands, heartbeat, results and artifacts",20,C["navy"],True)
    text(s,1650,1296,850,28,"LAN / VPN / central HTTPS  •  no SSH dependency",20,C["muted"],align="right")
    s.arrow([(1230,1250),(1230,1340)],C["purple"],2.5,head=10)
    s.arrow([(1410,1380),(1410,1250)],C["green"],2.5,head=10)
    s.line(426,1340,2106,1340,C["purple"],2.8)
    s.line(454,1380,2134,1380,C["green"],2.8)
    text(s,61,1326,340,28,"Command payloads ↓",20,C["purple"],True)
    text(s,61,1366,359,28,"Readbacks + evidence ↑",20,C["green"],True)
    for x,name in [(100,"PC i"),(940,"PC j"),(1780,"PC N")]:
        cx=x+340
        s.arrow([(cx-14,1340),(cx-14,1420)],C["purple"],2,head=8)
        s.arrow([(cx+14,1420),(cx+14,1380)],C["green"],2,head=8)
        s.circle(cx-14,1340,3.4,C["purple"]);s.circle(cx+14,1380,3.4,C["green"])
        s.rect(x,1420,680,218,C["ice"],C["teal"],radius=12,lw=1.7)
        monitor(s,x+21,1439,68,45,C["teal"])
        text(s,x+116,1440,541,37,f"{name} / instrument host",28,C["teal"],True)
        text(s,x+116,1480,541,27,"Local execution, vendor environment and data",20,C["muted"])
        s.rect(x+21,1524,638,49,C["white"],C["teal"],radius=7,lw=1.3)
        text(s,x+34,1538,612,31,"PC Agent → bounded drivers → device sessions",23,C["teal"],True,align="center")
        text(s,x+25,1596,630,29,"Watchdog • command journal • raw-data outbox",22,C["green"],align="center")
        s.arrow([(x+151,1638),(x+151,1690)],C["blue"],2.1,head=8)
        s.arrow([(x+169,1690),(x+169,1638)],C["blue"],2.1,head=8)
        text(s,x+205,1651,452,28,"SDK / VISA / USB / serial / LAN",20,C["blue"])
        s.rect(x+21,1690,638,90,C["bg"],C["blue"],radius=8,lw=1.4)
        text(s,x+38,1703,531,31,"Local instruments / controllers / detectors",22,C["blue"],True)
        text(s,x+38,1745,575,28,"Device 1 • Device 2 • … • Device m",20,C["muted"])
        s.rect(x+584,1707,55,28,C["pale"],C["blue"],radius=4)
        s.circle(x+624,1721,5,C["white"],C["blue"],1.1)
    text(s,802,1509,116,62,"…",48,C["faint"],True,align="center")
    text(s,1642,1509,116,62,"…",48,C["faint"],True,align="center")

    s.line(60,1809,2500,1809,C["line"],1.5)
    text(s,60,1823,1590,32,"Representative PCs: i, j, …, N. Any PC may host any registered instrument set.",23,C["navy"],True)
    text(s,1710,1824,790,31,"AI inference is central; PC-side LLMs are optional.",21,C["muted"],align="right")
    s.rect(60,1870,2440,39,C["ice"],radius=7)
    text(s,76,1878,2408,28,"Skills = reusable methods    |    MCP = tool interfaces    |    Coordinator = execution authority    |    PC Agent = local device I/O",22,C["teal"],True,align="center")
    return s


def verify(scene,pptx,pdf):
    errors=list(scene.errors);bounds=[]
    prs=Presentation(pptx);doc=pymupdf.open(pdf)
    assert len(prs.slides)==len(doc)==1
    ptext="\n".join(sh.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    pdftext=doc[0].get_text()
    assert ptext.count("ONE CENTRAL SERVER")==1
    assert ptext.count("PC Agent → bounded drivers → device sessions")==3
    for name in ["PC i","PC j","PC N"]:assert f"{name} / instrument host" in ptext
    for required in ["Skill library","Brainstorming","MCP host + clients","LabOS Control MCP","Local LLM • on this server"]:assert required in ptext
    for obsolete in ["Optical lab","Sample lab","Imaging lab","Motion lab","Electrical lab"]:assert obsolete not in ptext
    norm=lambda t:re.sub(r"\s+","",t)
    for oi,o in enumerate(scene.ops):
        if o["kind"]!="text":continue
        for j,line in enumerate(o["lines"]):
            if norm(line) not in norm(ptext):errors.append(f"Missing PPTX text: {line}")
            if norm(line) not in norm(pdftext):errors.append(f"Missing PDF text: {line}")
            tw=g.width(line,o["size"],o["bold"])
            xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
            yy=o["y"]+j*o["size"]*o["leading"]
            if line:bounds.append((oi,line,(xx,yy,xx+tw,yy+o["size"])))
    for i,a in enumerate(bounds):
        for b in bounds[i+1:]:
            if a[0]==b[0]:continue
            ax,ay,ar,ab=a[2];bx,by,br,bb=b[2]
            if min(ar,br)-max(ax,bx)>8 and min(ab,bb)-max(ay,by)>5:
                errors.append(f"Text overlap: {a[1]} / {b[1]}")
    for block in doc[0].get_text("dict")["blocks"]:
        for line in block.get("lines",[]):
            for span in line["spans"]:
                x0,y0,x1,y1=span["bbox"]
                if min(x0,y0)<0 or x1>W*.5+1 or y1>H*.5+1:errors.append(f"PDF text outside page: {span['text']}")
    assert not re.search(r"[\u3400-\u9fff]",ptext+prs.slides[0].notes_slide.notes_text_frame.text)
    assert not any(int(sh.shape_type)==13 for sh in prs.slides[0].shapes)
    if errors:raise AssertionError("\n".join(errors))
    return dict(status="PASS",pages=1,central_servers=1,representative_pc_symbols=3,fleet_size="N",fixed_pc_roles=False,
                skill_library=True,mcp_host_clients_connectors=True,pc_local_llm_required=False,language="English",raster_images_in_pptx=0,
                server_local_llm_fallback=True,
                native_objects=len(prs.slides[0].shapes),missing_text=0,out_of_page_text=0,text_overlaps=0)


def main():
    g.W=W;g.H=H;g.SCALE=.5;g.ROOT=ROOT;g.SLIDES.clear()
    scene=build();g.SLIDES.append(scene)
    if scene.errors:raise AssertionError("\n".join(scene.errors))
    title="LabOS System Architecture: Scientific Skills, MCP Tools and Distributed Execution"
    pptx=g.render_pptx("LabOS_System_Architecture_EN.pptx",title,"Agents, skills, MCP, model routing and generic instrument-host PCs",notes_source_label="Reference: ")
    pdf=g.render_pdf("LabOS_System_Architecture_EN.pdf",title)
    record=dict(title=scene.title,section=scene.section,note=scene.note,source=scene.source,ops=scene.ops)
    (ROOT/"architecture_scene.json").write_text(json.dumps(record,indent=2,ensure_ascii=False)+"\n")
    export_svg(record,ROOT/"LabOS_System_Architecture_EN.svg",(0,0,W,H))
    report=verify(scene,pptx,pdf)
    subprocess.run(["pdftoppm","-singlefile","-scale-to","5120","-png",str(pdf),str(ROOT/"LabOS_System_Architecture_EN")],check=True,capture_output=True)
    subprocess.run(["pdftoppm","-singlefile","-scale-to","2000","-png",str(pdf),str(ROOT/"Architecture_Preview")],check=True,capture_output=True)
    (ROOT/"QA_Report.json").write_text(json.dumps(report,indent=2)+"\n")
    files=sorted(p for p in ROOT.iterdir() if p.is_file() and p.name not in {"SHA256SUMS","LabOS_System_Architecture_EN_Bundle.zip"})
    (ROOT/"SHA256SUMS").write_text("".join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}\n" for p in files))
    bundle=ROOT/"LabOS_System_Architecture_EN_Bundle.zip"
    with zipfile.ZipFile(bundle,"w",compression=zipfile.ZIP_DEFLATED) as z:
        for p in files+[ROOT/"SHA256SUMS"]:z.write(p,p.name)
    with zipfile.ZipFile(bundle) as z:
        assert z.testzip() is None
        for p in files:assert z.read(p.name)==p.read_bytes()
    print(json.dumps(report,indent=2));print(f"Bundle integrity: PASS ({len(files)+1} files)")


if __name__=="__main__":main()
