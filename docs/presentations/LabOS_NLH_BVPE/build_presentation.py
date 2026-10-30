"""A shared vector scene generates an editable PPTX and a matching embedded-font PDF."""
from __future__ import annotations

import json
import math
import re
from pathlib import Path

from fontTools.ttLib import TTFont as FontInfo
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
W, H, SCALE = 1280, 720, .75
C = dict(navy="11233E", ink="23364D", muted="62758C", bg="F5F8FC", white="FFFFFF",
         line="D9E3EE", purple="7653CF", plum="F1ECFC", teal="168D9D", ice="E9F6F7",
         amber="BC7D1D", sand="FFF4DE", green="278565", mint="E9F6EF", red="B94C61",
         rose="FCEFF2", blue="3B6CB0", pale="EDF3FC", faint="AAB8C8")
FONT_CJK = "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"
FONT_LATIN = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
for n, p in (("CJK", FONT_CJK), ("Latin", FONT_LATIN), ("LatinB", FONT_BOLD)):
    pdfmetrics.registerFont(TTFont(n, p))
CMAP = FontInfo(FONT_CJK).getBestCmap()
LATIN_CMAP = FontInfo(FONT_LATIN).getBestCmap()


def font_for(char, bold=False):
    if ord(char) in CMAP:
        return "CJK"
    if ord(char) not in LATIN_CMAP:
        raise ValueError(f"missing glyph: {char!r} {ord(char):04x}")
    return "LatinB" if bold else "Latin"


def runs(text, bold=False):
    result = []
    for ch in text:
        font = font_for(ch, bold)
        if result and result[-1][0] == font:
            result[-1] = (font, result[-1][1] + ch)
        else:
            result.append((font, ch))
    return result


def width(text, size, bold=False):
    return sum(pdfmetrics.stringWidth(s, f, size) for f, s in runs(text, bold))


def wrap(text, available, size, bold=False):
    out = []
    for paragraph in text.split("\n"):
        if not paragraph:
            out.append("")
            continue
        line = ""
        for token in re.findall(r"[A-Za-z0-9_./:+-]+|.", paragraph):
            if width(line + token, size, bold) > available and line:
                out.append(line.rstrip())
                line = ""
            if width(token, size, bold) > available:
                for char in token:
                    if width(line + char, size, bold) > available:
                        out.append(line.rstrip())
                        line = ""
                    line += char
            else:
                line += token if line else token.lstrip()
        out.append(line.rstrip())
    return out


class Slide:
    def __init__(self, section, title, sub="", dark=False, note="", source=""):
        self.section, self.title, self.note, self.source = section, title, note, source
        self.ops, self.errors = [], []
        self.dark = dark
        self.rect(0, 0, W, H, C["navy"] if dark else C["bg"], border=None)
        if not dark:
            self.text(54, 31, 1080, 24, section.upper(), 13, C["teal"], True)
            self.text(54, 71, 1170, 57, title, 36, C["navy"], True)
            if sub:
                self.text(56, 131, 1165, 42, sub, 19, C["muted"])

    def rect(self, x, y, w, h, fill, border=None, radius=0, lw=1, dash=False):
        self.ops.append(dict(kind="rect", x=x, y=y, w=w, h=h, fill=fill, border=border,
                             radius=radius, lw=lw, dash=dash))

    def circle(self, x, y, r, fill, border=None, lw=1):
        self.ops.append(dict(kind="circle", x=x, y=y, r=r, fill=fill, border=border, lw=lw))

    def line(self, x1, y1, x2, y2, color=None, lw=2, dash=False):
        self.ops.append(dict(kind="line", x1=x1, y1=y1, x2=x2, y2=y2,
                             color=color or C["teal"], lw=lw, dash=dash))

    def poly(self, points, fill, border=None, lw=1):
        self.ops.append(dict(kind="poly", points=points, fill=fill, border=border, lw=lw))

    def arrow(self, points, color=None, lw=2.2, dash=False, head=9):
        color = color or C["teal"]
        for p1, p2 in zip(points, points[1:]):
            self.line(*p1, *p2, color, lw, dash)
        (x0, y0), (x1, y1) = points[-2:]
        angle = math.atan2(y1-y0, x1-x0)
        self.poly([(x1, y1), (x1-head*math.cos(angle)+head*.45*math.sin(angle),
                   y1-head*math.sin(angle)-head*.45*math.cos(angle)),
                  (x1-head*math.cos(angle)-head*.45*math.sin(angle),
                   y1-head*math.sin(angle)+head*.45*math.cos(angle))], color)

    def text(self, x, y, w, h, text, size=22, color=None, bold=False, align="left", leading=1.35):
        lines = wrap(text, w, size, bold)
        used = len(lines)*size*leading
        if used > h+2:
            self.errors.append(f"text overflow ({used:.0f}>{h}): {text[:65]}")
        if x < -1 or y < -1 or x+w > W+1 or y+h > H+1:
            self.errors.append(f"out of bounds: {text[:40]}")
        self.ops.append(dict(kind="text", x=x, y=y, w=w, h=h, lines=lines, size=size,
                             color=color or C["ink"], bold=bold, align=align, leading=leading))

    def pill(self, x, y, w, text, color=None, fill=None, size=15):
        self.rect(x, y, w, max(31,size*1.35+10), fill or C["ice"], radius=14)
        self.text(x+8, y+5, w-16, size*1.35+2, text, size, color or C["teal"], True, "center")

    def box(self, x, y, w, h, title, body="", accent=None, fill=None, size=21, title_size=24):
        accent = accent or C["teal"]
        self.rect(x, y, w, h, fill or C["white"], C["line"], radius=12)
        self.rect(x, y+17, 4, min(h-34, 33), accent)
        self.text(x+20, y+16, w-40, 39, title, title_size, accent, True)
        if body:
            self.text(x+20, y+62, w-40, h-74, body, size, C["ink"])

    def node(self, x, y, w, h, title, subtitle="", color=None, fill=None, size=22):
        color = color or C["teal"]
        self.rect(x, y, w, h, fill or C["white"], color, radius=10, lw=1.4)
        th = size*1.35
        top = y+(h-th)/2 if not subtitle else y+(8 if h<80 else 13)
        self.text(x+10, top, w-20, th+2, title, size, color, True, "center")
        if subtitle:
            sy=top+th+3
            self.text(x+12, sy, w-24, y+h-sy-5, subtitle, 14 if h<80 else 16, C["muted"], False, "center")

    def band(self, text, color=None, fill=None, y=618, size=21):
        self.rect(54, y, 1172, 54, fill or C["ice"], radius=9)
        self.text(73, y+12, 1134, 36, text, size, color or C["teal"], True)

    def footer(self, index, total):
        color = "B1C2DA" if self.dark else C["muted"]
        self.line(54, 687, 1226, 687, "31435B" if self.dark else C["line"], 1)
        self.text(55, 695, 1100, 19, self.source or "LabOS  |  非线性 Hall × BVPE  |  系统能力与应用场景", 11, color)
        self.text(1168, 693, 58, 21, f"{index:02d}/{total:02d}", 12, color, align="right")


SLIDES = []


def add(section, title, sub="", **kwargs):
    slide = Slide(section, title, sub, **kwargs)
    SLIDES.append(slide)
    return slide


def sample(s, x, y, w=230, h=110, label="同一样品"):
    s.rect(x, y, w, h, C["pale"], C["blue"], radius=12, lw=1.5)
    cy = y+h*.48
    s.rect(x+32, cy-12, w-64, 24, C["blue"], radius=4)
    for px in (x+w*.35, x+w*.65):
        s.line(px, cy-32, px, cy+32, C["blue"], 5)
        s.circle(px, cy-34, 4, C["amber"])
        s.circle(px, cy+34, 4, C["amber"])
    s.circle(x+24, cy, 5, C["amber"])
    s.circle(x+w-24, cy, 5, C["amber"])
    s.text(x+10, y+h-26, w-20, 25, label, 17, C["blue"], True, "center")


def timeline(s, x, y, items, w=1172, step_h=105):
    gap=18
    cw=(w-gap*(len(items)-1))/len(items)
    for i, (title, body) in enumerate(items):
        xx=x+i*(cw+gap)
        s.node(xx,y,cw,step_h,title,body,size=21)
        if i<len(items)-1:
            s.arrow([(xx+cw,y+step_h/2),(xx+cw+gap-2,y+step_h/2)],head=6)


def build_slides():
    s=add("LABOS", "LabOS", dark=True,
          note="这是一份系统能力与应用场景说明。我们用用户提供的两台 PC 和具体设备，说明 LabOS 如何组织非线性 Hall 与 BVPE 实验。图示展示的是系统架构、信号链与操作逻辑，不是材料测量结果报告。")
    s.pill(63,44,220,"MULTI-PC LAB SYSTEM", "BBA5FF", "263453",14)
    s.text(62,125,1000,110,"LabOS",78,C["white"],True)
    s.text(65,246,1130,115,"跨 PC · 跨仪器协同实验",48,C["white"],True)
    s.text(66,366,1090,65,"非线性 Hall × BVPE：一套装置、两条测量链、统一证据",26,"BDCEE4")
    s.arrow([(394,550),(560,550)],"5FC9D5",3)
    s.arrow([(720,550),(886,550)],"5FC9D5",3)
    s.node(66,495,328,112,"PC A · 光源与光谱","SuperK / IR spectrometer",C["teal"],C["ice"],24)
    s.node(560,495,160,112,"中央 Agent","计划 · 协调 · 复盘",C["purple"],C["plum"],21)
    s.node(886,495,328,112,"PC B · 环境与电学","attoDRY / lock-ins / gate / optics",C["teal"],C["ice"],24)
    s.text(67,641,1100,24,"2026.10  |  围绕科研问题组织设备，而不是围绕单台电脑组织实验",16,"BDCEE4")

    s=add("01 / RESEARCH SCENARIO","在同一套环境坐标下，关联电输运与光生响应",
          "共同坐标：样品、温度 T、磁场 B、栅压 Vg；两条测量链由中央统一组织。",
          note="非线性 Hall 与 BVPE 共享样品和环境条件，但激励、参考频率和读出链不同。LabOS 的价值是把这些条件与原始数据绑定起来，使暗态、光照态和回程对照能够比较。这里不预设某个具体材料的微观机制。")
    s.box(54,206,565,284,"非线性 Hall", "电激励：锁相输出 + 串联电阻\n输运通道：Vxx、Vxy\n测量维度：激励幅值、频率、T、B、Vg\n研究重点：二阶横向响应与对照",C["purple"],size=24,title_size=29)
    s.box(646,206,580,284,"BVPE / 光生响应", "光激励：SuperK 或 730 nm 激光\n调制参考：chopper → BVPE 锁相\n测量维度：波段、功率、偏振、T、B、Vg\n研究重点：光生信号与光学条件关联",C["teal"],size=24,title_size=29)
    for x,w,text in ((92,245,"同一份样品记录"),(365,245,"同一套环境读回"),(638,245,"可比较的测量阶段"),(911,275,"可追溯的数据与决策")):
        s.pill(x,535,w,text,size=18)
    s.band("中央 Agent 把科研目标拆成可执行阶段，再把不同 PC 的结果汇成同一条证据链。",size=22)

    s=add("02 / SYSTEM ARCHITECTURE","中央做编排，本地管设备，数据反馈形成闭环",
          "总体分层借鉴 QUANTA 的框架表达，按 LabOS 与这套双 PC 装置重新绘制。",
          source="[1] QUANTA, arXiv:2610.38850, Fig. 1；本图为 LabOS 场景原创重绘。",
          note="从上到下分成操作者、科学规划、确定性执行以及本地设备。中央 Agent 利用资料和历史数据形成阶段计划；确定性执行层负责身份、资源归属、参数范围和状态检查。PC A 和 PC B 通过出站网络访问中央；PC 本地持有厂商连接。数据和异常沿反馈路径返回中央，下一步建议仍可以由操作者审核。")
    s.arrow([(632,224),(632,244)],C["amber"])
    s.arrow([(632,322),(632,353)],C["purple"])
    s.arrow([(530,437),(530,446),(328,446),(328,456)])
    s.arrow([(749,437),(749,446),(941,446),(941,456)])
    s.arrow([(235,610),(30,610),(30,285),(277,285)],C["green"],2,dash=True)
    s.arrow([(1050,610),(1250,610),(1250,285),(1001,285)],C["green"],2,dash=True)
    s.arrow([(199,345),(199,310),(278,310)],C["purple"],2)
    s.arrow([(856,395),(941,395)],C["green"],2)
    s.arrow([(1100,345),(1100,311),(1001,311)],C["green"],2)
    s.node(382,183,500,43,"研究者：目标、样品、允许边界与物理操作",color=C["amber"],fill=C["sand"],size=19)
    s.node(278,244,723,80,"中央 Agent · 科研规划与解释","任务分解  /  资料检索  /  数据审阅  /  诊断与下一步建议",C["purple"],C["plum"],25)
    s.node(54,345,286,73,"知识与实验记录","手册 · 旧代码 · 协议 · 历史运行",C["purple"],C["white"],21)
    s.node(423,353,433,84,"确定性执行协调器","参数校验 · 设备独占 · 阶段屏障 · 命令日志",C["teal"],C["ice"],24)
    s.node(941,345,285,73,"统一数据与证据","原始值 · 读回 · 质量标记 · 复盘",C["green"],C["white"],21)
    s.box(77,457,500,168,"PC A Agent · 光学资源", "SuperK 超连续光源\n红外光谱仪：光谱与参考采集",C["teal"],size=23)
    s.box(625,457,579,168,"PC B Agent · 环境、电学与调制", "attoDRY · Vxx / Vxy 锁相 · BVPE 锁相\nchopper · PD · 2400 · 730 nm / CLD1000\n电控偏振片 / 四分之一波片",C["teal"],size=21)
    s.text(43,438,290,23,"测量结果 / 异常 / 质量反馈",15,C["green"])

    s=add("03 / DEVICE OWNERSHIP","设备归属清楚，跨 PC 协作就有稳定边界",
          "每台设备由所属 PC 的 Agent 持有连接；中央调用可描述、可组合的仪器能力。",
          note="这一页完整对应用户给出的设备清单。PC A 负责 SuperK 与红外光谱，PC B 负责环境、电输运、栅压和调制光学。注意输运一共有两台锁相，另有一台锁相用于 chopper 参考的光生信号，总共三台。PD 测光功率，CLD1000 调激光驱动电流，功率由标定与 PD 读回关联。")
    s.rect(54,191,1172,418,C["white"],C["line"],radius=12)
    cols=[76,189,638]
    for x,txt in zip(cols,["归属","设备 / 子系统","在实验中的作用"]):
        s.text(x,204,104 if x==76 else 540,29,txt,20,C["muted"],True)
    rows=[("PC A","SuperK 超连续光源","宽谱 / 选定波段的光激励"),
          ("PC A","红外光谱仪","入射谱 / 参考谱 / 光谱诊断"),
          ("PC B","attoDRY","样品温度 T 与磁场 B 控制"),
          ("PC B","输运锁相 × 2","分别读 Vxx、Vxy；配置 1f / 2f 与 X/Y 相位"),
          ("PC B","chopper + BVPE 锁相","光调制与同步光生电流 / 电压读出"),
          ("PC B","PD + 2400","光功率读回；栅压 Vg 与栅漏电流 Ig"),
          ("PC B","730 nm 激光器 + CLD1000","激光照明；驱动电流与光功率标定"),
          ("PC B","电控偏振片 / 四分之一波片","线偏振角 / 偏振态扫描")]
    for i,(pc,device,role) in enumerate(rows):
        yy=242+i*43
        if i%2==0:s.rect(65,yy-3,1150,42,C["bg"])
        s.text(cols[0],yy,100,32,pc,19,C["teal"],True)
        s.text(cols[1],yy,430,32,device,19,C["ink"],True)
        s.text(cols[2],yy,560,33,role,19,C["ink"])
    s.band("同一台设备只归一个 Agent 管理；同一个实验可以组合任意多个 PC 的能力。",size=22)

    s=add("04 / PHYSICAL SIGNAL MAP","软件连成一张网，物理信号仍沿各自链路流动",
          "青色：光路与测量信号；金色虚线：物理参考；紫色：电激励与控制。",
          note="这张图说明两台 PC 的分工并不限制物理光路。PC A 的 SuperK 可以把光送到 PC B 管理的 attoDRY 样品，红外光谱仪可对参考分支进行诊断。730 nm 是另一条可配置光源。CLD1000 是电流控制，不是串在光路里的光学元件。chopper 参考与电激励参考是物理线。共同电极在不同测量阶段按接线配置使用，图示没有假定安装自动切换矩阵。")
    s.rect(54,190,306,198,C["ice"],radius=12)
    s.text(72,201,265,32,"PC A · 光源 / 光谱",22,C["teal"],True)
    s.node(73,247,265,54,"SuperK + 波段配置",color=C["teal"],size=20)
    s.node(73,320,265,50,"IR 光谱仪 · 参考谱",color=C["teal"],size=19)
    s.rect(54,405,306,180,C["plum"],radius=12)
    s.text(72,416,265,32,"PC B · 激光光源",22,C["purple"],True)
    s.node(73,462,265,48,"730 nm 激光器",color=C["purple"],size=21)
    s.node(73,527,265,42,"CLD1000 · 驱动电流",color=C["purple"],size=18)
    s.arrow([(205,527),(205,510)],C["purple"],2)
    s.arrow([(338,274),(386,274),(386,369),(431,369)],C["teal"],3)
    s.arrow([(338,485),(386,485),(386,369)],C["teal"],3)
    s.node(431,337,151,65,"chopper","光强调制",C["teal"],size=20)
    s.node(615,337,154,65,"偏振模块","线偏振片 / λ/4",C["teal"],size=20)
    s.arrow([(582,369),(615,369)],lw=3)
    s.arrow([(769,369),(821,369)],lw=3)
    s.rect(821,288,207,192,C["pale"],C["blue"],radius=13)
    s.text(840,303,168,30,"attoDRY · T / B",21,C["blue"],True,"center")
    sample(s,833,355,183,100,"样品 + gate")
    s.arrow([(790,369),(790,249),(714,249)],C["teal"],2)
    s.node(549,218,165,59,"PD · 功率读回",color=C["teal"],size=19)
    s.arrow([(386,345),(338,345)],C["teal"],1.7,dash=True)
    s.text(382,207,155,60,"光源与参考光路\n由实验配置选择",17,C["muted"])
    s.node(1062,250,164,61,"锁相 · Vxx",color=C["purple"],size=21)
    s.node(1062,340,164,61,"锁相 · Vxy",color=C["purple"],size=21)
    s.node(1062,461,164,69,"BVPE 锁相","光生信号读出",C["teal"],size=20)
    s.arrow([(1028,369),(1043,369),(1043,281),(1062,281)],C["purple"],2)
    s.arrow([(1028,390),(1048,390),(1048,370),(1062,370)],C["purple"],2)
    s.arrow([(1028,417),(1040,417),(1040,496),(1062,496)],C["teal"],2)
    s.arrow([(507,402),(507,567),(1137,567),(1137,530)],C["amber"],2,dash=True)
    s.text(600,575,466,27,"chopper TTL / REF → BVPE 锁相",18,C["amber"],True)
    s.node(814,510,194,44,"2400 → Vg / Ig",color=C["purple"],size=19)
    s.arrow([(912,510),(912,480)],C["purple"],2)
    s.band("Hall 与 BVPE 按测量阶段使用接线配置；电激励与参考连接见后续测量链路页。",size=21)

    s=add("05 / DISTRIBUTED CONTROL","中央派发实验阶段，PC Agent 执行本地动作",
          "Windows PC 主动连接中央：跨局域网或专用网络组合仪器，无需以 SSH 作为运行前提。",
          note="网络层传递的是结构化命令和可追溯的数据，不是高频电信号。中央运行器持久化命令、维护设备资源独占和实验阶段。Agent 本地调用已封装的驱动，保存原始数据和执行日志，再校验上传。心跳、会话和命令 ID 支持断线后的状态核对；已经发出但结果不明的物理动作不能盲目重放。")
    s.node(473,194,334,100,"中央服务 / Agent","实验计划 · 资源锁 · 命令与证据",C["purple"],C["plum"],27)
    for xx,title,body in [(75,"PC A Agent","光源 / 光谱驱动\n本地数据 / 执行日志"),(843,"PC B Agent","环境 / 电学 / 光学驱动\n本地数据 / 执行日志")]:
        s.node(xx,393,362,125,title,body,C["teal"],C["white"],27)
    s.arrow([(472,244),(254,244),(254,393)],C["purple"],2.5)
    s.arrow([(807,244),(1024,244),(1024,393)],C["purple"],2.5)
    s.arrow([(437,453),(530,453),(530,294)],C["green"],2.5)
    s.arrow([(843,453),(750,453),(750,294)],C["green"],2.5)
    s.text(83,284,285,77,"PC 主动轮询\n接收批准的阶段命令",23,C["purple"],True)
    s.text(926,284,295,77,"心跳 / 状态 / 原始数据\n经 HTTP(S) 返回中央",23,C["green"],True)
    s.node(510,516,260,66,"共享实验身份","campaign / run / point",C["blue"],C["pale"],21)
    s.band("科研编排跨 PC；设备连接留在本地；原始数据先落盘，再可靠同步。",size=23)

    s=add("06 / TIMING AND REFERENCES","软件屏障协调实验点，物理参考决定测量相位",
          "共同 point_id 绑定条件与采集窗口；网络到达时间不承担锁相同步。",
          note="中央先向 PC A 设置光学条件，同时让 PC B 设置温度、磁场和栅压；两端提交新鲜读回，进入该实验点的准备屏障。完成预先选择的稳定性判断后，中央发出采集阶段。各设备保留采集起止时间，因此跨 PC 数据是有共同身份和时间窗口的对应数据，并不默认硬件同时采样。电输运锁相共享物理 fI 参考，BVPE 锁相使用 chopper 的 fchop 参考。")
    tracks=[("中央",C["purple"]),("PC A",C["teal"]),("PC B",C["teal"]),("物理参考",C["amber"])]
    for i,(label,color) in enumerate(tracks):
        yy=218+i*91
        s.text(56,yy+16,142,42,label,23,color,True)
        s.line(209,yy+40,1215,yy+40,C["line"],2)
    s.node(212,220,196,54,"分发 point_id",color=C["purple"],fill=C["plum"],size=20)
    s.node(476,220,221,54,"汇总两端 READY",color=C["purple"],fill=C["plum"],size=20)
    s.node(786,220,220,54,"启动采集阶段",color=C["purple"],fill=C["plum"],size=20)
    s.node(1060,220,153,54,"归档与评估",color=C["purple"],fill=C["plum"],size=19)
    s.node(249,311,247,54,"设置光源 / 光谱条件",size=20)
    s.node(527,311,170,54,"读回并 READY",size=19)
    s.node(786,311,285,54,"参考光谱 / 光学采集",size=20)
    s.node(249,402,247,54,"设置 T / B / Vg 并稳定",size=20)
    s.node(527,402,170,54,"读回并 READY",size=19)
    s.node(786,402,285,54,"输运 / BVPE / PD 采集",size=20)
    s.line(741,196,741,477,C["purple"],2,True)
    s.text(680,476,133,29,"阶段屏障",19,C["purple"],True,"center")
    s.arrow([(245,533),(1193,533)],C["amber"],3)
    s.text(257,548,930,37,"电输运：fI 参考线     |     光生测量：fchop 参考线",23,C["amber"],True)
    s.band("每个采集窗口记录实际条件、时间戳和锁定状态；需要严格同时采样时组合硬件触发。",size=21)

    s=add("07 / PREPARATION","实验前，把目标、设备与参数组织成一份计划",
          "设备清单是可复用配置；样品、接线与本次扫描条件形成实验准备记录。",
          note="首次接入时登记 PC 与设备归属、仪器连接、驱动及可用能力，之后复用。每次实验填写样品、测量接线、研究问题、扫描变量与允许范围。中央将其组织为多阶段计划，准备记录包含光学标定、激励标定、功率、栅漏电流、锁相参考和收尾策略。物理操作由操作者完成和记录。")
    s.box(54,200,364,349,"设备配置 · 可复用", "PC A / PC B 与设备归属\n仪器地址、驱动与能力\n锁相输入 / 输出 / 参考\n光源、偏振与功率标定\n现用脚本与操作协议",C["teal"],size=23)
    s.box(457,200,364,349,"实验配置 · 每次记录", "样品 ID 与接线版本\n非线性 Hall / BVPE 模式\nT、B、Vg 与扫描顺序\n激励 / 光功率 / 偏振参数\n对照、重复与结束状态",C["purple"],size=23)
    s.box(860,200,366,349,"中央输出 · 可执行计划", "有序阶段与依赖关系\n各 PC 的设备动作\n读回、稳定与采集条件\n数据质量与异常处理\n需要操作者执行的步骤",C["amber"],size=23)
    s.band("描述“要研究什么”，再核对“设备怎么配合、每一步以什么证据判断完成”。",size=23)

    s=add("08 / NONLINEAR HALL — SIGNAL CHAIN","非线性 Hall：两个空间通道，共用电激励参考",
          "Vxx 主锁相同时提供激励与 1f 读出；另一台测 Vxy(2f)，后续阶段扩展其他谐波。",
          source="[2] Sodemann & Fu, PRL 115, 216806 (2015)；[3] Ma et al., Nature 565, 337–342 (2019)。",
          note="用户确认交流激励由锁相输出和串联电阻提供。图中左侧的正弦输出与右侧的 Vxx 读出属于同一台主锁相的两个功能，不是额外增加一台输运锁相。主锁相使用内部参考，同时把 fI 的物理参考送给测 Vxy 的锁相，并在后者配置二次谐波。实际样品电流通过串联电阻压降等标定方法获得。Vxx/Vxy 是空间电压通道，各锁相还返回 X/Y 相位分量。其他谐波可按阶段切换，或组合多解调仪器。")
    s.node(64,236,249,106,"Vxx 锁相输出","电激励频率 fI（内参）",C["purple"],C["plum"],24)
    s.node(359,249,171,80,"串联电阻 Rs","标定实际 Iω",C["purple"],size=22)
    sample(s,600,242,243,130,"attoDRY 中的样品")
    s.arrow([(313,289),(359,289)],C["purple"],3)
    s.arrow([(530,289),(600,289)],C["purple"],3)
    s.arrow([(843,354),(876,354),(876,431),(113,431),(113,342)],C["purple"],2)
    s.text(297,436,474,26,"电激励回路返回",17,C["purple"],True,"center")
    s.arrow([(843,276),(898,276),(898,245),(962,245)],C["blue"],2)
    s.arrow([(843,322),(898,322),(898,371),(962,371)],C["blue"],2)
    s.node(962,204,253,86,"输运锁相 · Vxx","内参 1f → Vxx(ω)",C["blue"],C["pale"],23)
    s.node(962,329,253,86,"输运锁相 · Vxy","n = 2 → Vxy(2ω)",C["blue"],C["pale"],23)
    s.arrow([(189,342),(189,470),(1090,470),(1090,415)],C["amber"],2.5,True)
    s.text(344,480,755,41,"主锁相内参 fI；物理 REF → Vxy 锁相",23,C["amber"],True,"center")
    s.pill(64,535,425,"空间通道：Vxx / Vxy",C["blue"],C["pale"],21)
    s.pill(538,535,677,"每台锁相：同时保存 X、Y、量程与参考状态",C["blue"],C["pale"],21)
    s.band("Vxy(2ω) 与 Iω² 的关系用于检验二阶响应；微观机制由对照证据进一步区分。",size=21)

    s=add("09 / NONLINEAR HALL — PROTOCOL","让二次响应与环境、激励和对照共同进入证据链",
          "中央把 T / B / Vg 的外层扫描，与各点的激励测试和锁相采集组合起来。",
          note="先做暗态基线和低激励检查，再进行幅值扫描与返回参考点。每个环境点可配置重复、反向磁场和频率对照。Vxy 的二次激励依赖并不能自动证明 Berry 曲率偶极机制；热电二次响应、接触非线性和纵向串扰需要相应控制。磁场奇偶分析也应按候选模型区分，不能套用线性 Hall 的单一规则。")
    timeline(s,54,206,[("环境准备","T / B / Vg → 读回"),("暗态基线","1f 基线 / 量程"),("幅值扫描","Iω 列表 + 2f"),("回程对照","参考点 / 频率 / 重复"),("评估与下一步","质量 + 证据 + 建议")],step_h=116)
    s.box(54,373,565,204,"中央组织的扫描空间", "外层：T × B × Vg\n内层：Iω × 频率 × 谐波 × 重复\n每点绑定实际电流与采集时间窗",C["purple"],size=25)
    s.box(650,373,576,204,"保留可区分机制的对照", "幅值依赖与返回参考\n频率 / 相位 / 接线与串扰检查\n热效应、漏电与长期漂移记录",C["teal"],size=25)
    s.band("阶段之间可以根据数据质量和对照结果，提出新的测量条件与范围。",size=23)

    s=add("10 / BVPE — SIGNAL CHAIN","BVPE：光学条件可编排，光生信号锁定 chopper",
          "730 nm / SuperK → 调制 → 偏振 → 样品；独立 PD 把光功率纳入每个测量点。",
          source="[4] Osterhoudt et al., Nature Materials 18, 471–475 (2019)；测量链为 LabOS 场景示意。",
          note="激光经过 chopper 与偏振模块照射样品。chopper 的物理参考送到第三台锁相，用来在 fchop 读取光生响应；光电流可通过仪器电流输入或前置电流放大器读出，光生电压则使用电压输入。PD 通过响应度与光路传输标定关联样品面功率。光学二阶响应的光频与 chopper 的低频包络是不同概念，因此 BVPE 不等于必须在 2fchop 读数。源漏零外加偏压可以与有限栅压同时存在。")
    s.node(54,217,226,84,"SuperK / 730 nm","选择光源与激发波段",C["teal"],C["ice"],22)
    s.node(329,217,180,84,"chopper","调制 fchop",C["teal"],size=23)
    s.node(558,217,196,84,"偏振模块","角度 θ / φ",C["teal"],size=23)
    sample(s,818,202,230,137,"样品 · T / B / Vg")
    for a,b in [(280,329),(509,558),(754,818)]:s.arrow([(a,259),(b,259)],lw=3)
    s.node(818,417,230,88,"BVPE 锁相","光生电流 / 电压 · X/Y",C["blue"],C["pale"],24)
    s.arrow([(932,339),(932,417)],C["blue"],2.5)
    s.arrow([(419,301),(419,462),(818,462)],C["amber"],2.5,True)
    s.text(452,475,335,39,"物理 REF：fchop",23,C["amber"],True)
    s.node(565,369,188,83,"PD 功率监测","Psample 的标定映射",C["green"],C["mint"],22)
    s.arrow([(784,259),(784,410),(753,410)],C["green"],2)
    s.node(54,378,340,92,"CLD1000 → 730 nm","驱动设置 → PD 读回 → 功率映射",C["purple"],C["plum"],21)
    s.arrow([(164,378),(164,301)],C["purple"],2)
    s.text(71,534,1130,48,"BVPE 测量：源漏零外加偏压；栅压作为独立调控轴；光生信号由对照与标定支撑。",22,C["ink"],True)
    s.band("电学解调频率 fchop 对应光强调制包络，区别于激发光频与 Hall 的 2fI。",size=23)

    s=add("11 / BVPE — PROTOCOL","功率、偏振与光谱组合成可复用的测量阶段",
          "中央将 PC A 的光源 / 光谱，与 PC B 的环境、偏振、调制和电读出配对。",
          note="先采集暗态与参考光谱，再建立光功率标定。730 nm 支持固定波长下的功率和偏振扫描；SuperK 根据所配选波能力进行波段测量，也可把完整入射谱作为宽谱条件记录。偏振片测线偏振角依赖，四分之一波片配合已知线偏振输入和延迟量标定生成不同偏振态。红外光谱仪按其响应波段参与参考，不把它默认当成 730 nm 功率计。光生响应归一化应说明功率、放大器增益以及调制幅度的定义。")
    timeline(s,54,206,[("暗态与参考","暗背景 / 光谱 / 零偏置"),("功率标定","驱动 → PD → Psample"),("偏振扫描","线偏振角 / λ/4 角"),("多维测量","光谱 × T × B × Vg"),("返回参考","暗态 / 参考功率 / 重复")],step_h=116)
    s.box(54,374,365,205,"730 nm · 固定波长", "CLD1000 驱动设置\nPD 监测与功率归一\n功率 / 偏振态依赖",C["purple"],size=24)
    s.box(457,374,365,205,"SuperK · 波段配置", "设置激发波段\nIR 光谱记录入射条件\n光谱响应与环境扫描",C["teal"],size=24)
    s.box(861,374,365,205,"偏振模块 · 可切换", "线偏振片：θ 扫描\nλ/4 波片：φ 与偏振态\n保存角度原点与标定",C["blue"],size=24)
    s.band("中央比较的是带有实际光功率、偏振配置与入射谱证据的光生响应。",size=23)

    s=add("12 / JOINT CAMPAIGN","一个研究任务，串起暗态输运与光生测量",
          "各阶段共享样品与环境坐标；以 campaign / run / point 关联结果和条件。",
          note="这是系统能力展示的核心案例。研究者问同一样品在温度、磁场和栅压改变时，二阶横向电响应与光生响应如何变化。中央组织暗态输运、光学参考、光生测量、回程对照四类阶段。换光源或重新接线等物理操作进入操作者步骤。可以比较共同坐标下的信号趋势，但共享坐标不意味着两个信号已经有相同微观起源。阶段记录中的 campaign 和 point 是这里展示的共同身份模型。")
    xs=[54,356,658,960]
    entries=[("A / 暗态输运","PC B","T / B / Vg → Vxx / Vxy",C["purple"]),
             ("B / 光学参考","PC A + PC B","光谱 + PD → 光学条件",C["teal"]),
             ("C / 光生响应","PC A + PC B","光源 / 偏振 → BVPE 锁相",C["blue"]),
             ("D / 返回对照","PC B + PC A","暗态 / 功率 / 参考点复测",C["green"])]
    for i,(title,pc,body,color) in enumerate(entries):
        s.box(xs[i],223,266,212,title,pc+"\n\n"+body,color,size=22,title_size=22)
        if i<3:s.arrow([(xs[i]+266,327),(xs[i+1]-5,327)],color,lw=2.4)
    s.rect(54,483,1172,85,C["white"],C["line"],radius=12)
    s.text(73,499,1132,49,"共同坐标：sample_id + Tread + Bread + Vgread + 采集时间窗 + 接线版本",25,C["navy"],True,"center")
    s.band("先把条件和证据对齐，再讨论电响应、光响应及其环境依赖之间的关系。",size=23)

    s=add("13 / QUALITY AND PROVENANCE","每一个点，都能解释“测了什么、条件如何、能否使用”",
          "请求值与实际读回分开保存，原始信号与质量判断分开保存。",
          note="记录包括计划、配置指纹、共同实验身份以及每次采集的起止时间。输运锁相保存 X/Y 与谐波，光生测量保存 chopper 参考和前放增益，PD 与光谱记录关联光学条件。过载、失锁或不满足预定稳定标准的数据进入拒绝或隔离状态，原始值仍保留。故障后的输出处理依照设备和实验策略：关掉光或激励不等于磁场、温度、栅压全部瞬间归零。")
    s.box(54,203,355,347,"共同身份与条件", "sample / campaign / point\n计划与配置哈希\n请求值 + 实际读回\n采集起止时间\n光源 / 偏振 / 接线版本",C["purple"],size=23)
    s.box(443,203,355,347,"原始信号与仪器状态", "Vxx、Vxy 的 X/Y 与谐波\nBVPE 原始信号与增益\nPD 功率 / 入射光谱\n过载、参考锁定与量程\n本地原始文件与摘要",C["teal"],size=23)
    s.box(831,203,395,347,"质量判定与实验复盘", "合格：进入后续分析\n拒绝：保留原始数据\n历史：注明状态与时间\n记录关闭输出后的读回\n形成有依据的后续建议",C["green"],size=23)
    s.band("统一数据不是只拼接数值，而是把条件、来源、质量与解释边界一起保存。",size=23)

    s=add("14 / USER EXPERIENCE","研究者描述目标，中央把跨 PC 流程组织起来",
          "设备配置可复用；每次实验聚焦样品、变量、对照与允许范围。",
          note="研究者可以用自然语言描述任务，然后在中央核对形成的计划。示例中的温度、磁场、栅压和角度不指定具体数值，因为这些都属于本次实验配置。每台 Windows PC 运行本地 Agent；研究者不必轮流远程登录各台电脑去运行脚本。物理接线、样品更换和关键阶段批准仍有明确的人机交接。")
    s.rect(54,198,1172,112,C["plum"],radius=12)
    s.text(79,217,1120,77,"“对同一样品，按 T、B、Vg 扫描暗态非线性 Hall；\n再在对应环境点做 730 nm 的功率与偏振依赖，并加入返回参考。”",27,C["purple"],True)
    timeline(s,54,364,[("描述目标","样品 / 科学问题"),("核对准备","设备 / 接线 / 参数"),("审阅计划","阶段 / 范围 / 对照"),("协同执行","状态 / 数据 / 异常"),("查看复盘","证据 / 解释 / 建议")],step_h=128)
    s.text(74,541,1130,75,"PC 端配置：PC 身份 + 仪器连接 + 驱动 / 现用代码 + 数据目录\n中央端配置：实验目标 + 样品条件 + 扫描范围 + 阶段与结束策略",23,C["ink"])
    s.band("研究者把精力放在实验设计与判断；中央与本地 Agent 负责组织、执行和留痕。",size=22)

    s=add("15 / REUSE EXISTING CODE","把已有测试脚本，封装成可以跨实验复用的能力",
          "本地控制环境继续发挥作用；中央通过统一操作描述组合这些能力。",
          note="用户原有设备测试代码非常有价值，可以保留厂商环境和已验证的调用。我们先识别连接、状态、设置、采集和结束步骤，再封装成有限操作。对外能力只暴露参数和结构化结果，连接和设备细节留在本地。已有脚本还可作为可检索资料，为中央规划提供依据；作为资料导入与作为驱动执行是两条不同路径。")
    s.box(54,205,334,303,"现用仪器测试代码", "连接和设备识别\n读状态 / 设置参数\n采集和原始保存\n异常处理 / 关闭输出",C["purple"],size=24)
    s.box(473,205,334,303,"本地仪器能力封装", "connect / status\nset / settle\nacquire / save\nstop / verify",C["teal"],size=25)
    s.box(892,205,334,303,"中央可组合任务", "环境准备阶段\nHall 测量阶段\n光学标定阶段\nBVPE 与返回对照",C["blue"],size=24)
    s.arrow([(388,350),(473,350)],C["teal"],3)
    s.arrow([(807,350),(892,350)],C["teal"],3)
    s.text(73,546,1130,50,"协议、代码版本与测试记录也进入资料库：让“这台设备如何可靠测量”成为可积累的知识。",23,C["ink"])
    s.band("同一套设备能力可以参与多个实验，原有控制代码可以逐步沉淀为统一接口。",size=22)

    s=add("16 / SCALE OUT","从两台 PC 扩展到更多实验台，沿用相同协同方式",
          "扩展的单位是“PC Agent + 设备能力 + 实验记录”，而不只是远程桌面数量。",
          note="LabOS 的扩展不要求所有仪器搬到一台电脑。新增实验台向中央登记资源、能力和控制归属，就可以被纳入后续计划。共同资源仍要独占；共享的光源、样品与物理路径用配置和阶段依赖表达。网络层的统一不会替代每种仪器所需的物理参考和触发。多实验室共享时还可以按网络部署和授权边界配置中央入口。")
    s.node(460,203,360,92,"中央 Agent / 统一实验记录","资源目录 · 阶段编排 · 质量与复盘",C["purple"],C["plum"],26)
    coords=[(68,"PC A","光源与光谱"),(374,"PC B","低温输运与 BVPE"),(680,"PC C","扫描 / 成像能力"),(986,"PC N","另一实验台 / 新仪器")]
    for xx,title,body in coords:
        s.arrow([(640,295),(640,337),(xx+115,337),(xx+115,384)],lw=2.1)
        s.node(xx,384,225,126,title,body,C["teal"],C["white"],27)
    s.pill(225,552,244,"统一资源与能力目录",size=18)
    s.pill(516,552,244,"跨 PC 阶段依赖",size=18)
    s.pill(807,552,244,"统一数据与证据",size=18)
    s.band("新增设备接入本地 Agent，新增实验复用已有能力，中央继续围绕研究问题编排。",size=22)

    s=add("LABOS / TAKEAWAY","一份计划，协同多台 PC；一条证据链，连接多种响应",dark=True,
          note="最后总结系统价值。它能够组织电激励和光激励两条链，组合跨 PC 的光源、低温环境和读出资源，给每个测量点建立条件与来源清楚的记录，并把数据质量和下一步建议接到实验流程里。这里展示的是装置场景和系统能力，而不是对具体样品已有机制结论的宣告。")
    s.pill(62,45,196,"LABOS CAPABILITIES","BBA5FF","263453",14)
    s.text(62,126,1140,138,"一份计划，协同多台 PC\n一条证据链，连接多种响应",45,C["white"],True)
    items=[("跨 PC 协同","光源、环境、锁相与偏振资源统一组织"),
           ("围绕科学问题","Hall 与 BVPE 的阶段、对照和变量可组合"),
           ("结果可追溯","条件、原始值、质量与解释共同保存")]
    for i,(title,body) in enumerate(items):
        yy=343+i*89
        s.circle(77,yy+17,9,"5FC9D5")
        s.text(105,yy,273,43,title,29,"71D8DF",True)
        s.text(405,yy+3,805,45,body,25,"BDCEE4")
    s.text(64,640,1130,27,"非线性 Hall × BVPE  |  电激励与光激励，共享同一套实验组织与证据语言",20,"BDCEE4")

    s=add("APPENDIX / SOURCES","参考与图示说明",
          "框架表达参考论文；全部系统、光路、电路与时序图均按本场景重新绘制。",
          note="资料来源分为三类：用户确认的装置清单与激励方式，LabOS 当前仓库中的系统设计与实现，以及公开科学论文。QUANTA 主要用于参考分层架构表达；非线性 Hall 与 BVPE 文献用于说明激励、谐波和对照逻辑。演示不使用文献中的材料曲线冒充本系统实测结果。")
    refs=[("[1] 系统框架", "Rha et al., Autonomous Quantum Measurement with Agentic AI, 2026.","https://arxiv.org/abs/2610.38850"),
          ("[2] 非线性 Hall 理论", "Sodemann & Fu, Physical Review Letters 115, 216806, 2015.","https://doi.org/10.1103/PhysRevLett.115.216806"),
          ("[3] 非线性 Hall 实验", "Ma et al., Nature 565, 337–342, 2019.","https://doi.org/10.1038/s41586-018-0807-6"),
          ("[4] BVPE 实验", "Osterhoudt et al., Nature Materials 18, 471–475, 2019.","https://doi.org/10.1038/s41563-019-0297-4")]
    for i,(title,body,url) in enumerate(refs):
        yy=194+i*94
        s.text(65,yy,1120,30,title,21,C["teal"],True)
        s.text(65,yy+31,1120,29,body,20,C["ink"])
        s.text(65,yy+60,1120,26,url,17,C["blue"])
    s.band("设备清单与激励方式：用户提供并确认。图中的设备接口、阶段和参数按实验配置组织。",size=20)


def rgb(value):
    return tuple(int(value[i:i+2],16)/255 for i in (0,2,4))


def pval(value):
    return Pt(value*SCALE)


def ppt_color(value):
    return RGBColor.from_string(value)


def style_shape(shape, fill=None, border=None, lw=1, dash=False):
    if hasattr(shape,"fill"):
        if fill:
            shape.fill.solid(); shape.fill.fore_color.rgb=ppt_color(fill)
        else:
            shape.fill.background()
    if border:
        shape.line.color.rgb=ppt_color(border); shape.line.width=pval(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    else:
        shape.line.fill.background()


def render_pptx(filename="LabOS_NLH_BVPE_System.pptx", title=None, subject=None, notes_source_label="来源："):
    prs=Presentation()
    prs.slide_width=pval(W); prs.slide_height=pval(H)
    prs.core_properties.title=title or "LabOS：跨 PC 协同的非线性 Hall 与 BVPE 实验系统"
    prs.core_properties.subject=subject or "系统能力、双 PC 装置与测量流程"
    prs.core_properties.author="LabOS"
    for scene in SLIDES:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        for o in scene.ops:
            k=o["kind"]
            if k=="rect":
                kind=MSO_SHAPE.ROUNDED_RECTANGLE if o["radius"] else MSO_SHAPE.RECTANGLE
                sh=slide.shapes.add_shape(kind,pval(o["x"]),pval(o["y"]),pval(o["w"]),pval(o["h"]))
                if o["radius"]:
                    sh.adjustments[0]=min(.5,o["radius"]/min(o["w"],o["h"]))
                style_shape(sh,o["fill"],o["border"],o["lw"],o["dash"])
            elif k=="circle":
                sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,pval(o["x"]-o["r"]),pval(o["y"]-o["r"]),pval(o["r"]*2),pval(o["r"]*2))
                style_shape(sh,o["fill"],o["border"],o["lw"])
            elif k=="poly":
                pts=o["points"]
                b=slide.shapes.build_freeform(pts[0][0]*SCALE,pts[0][1]*SCALE,scale=12700)
                b.add_line_segments([(x*SCALE,y*SCALE) for x,y in pts[1:]],close=True)
                sh=b.convert_to_shape(); style_shape(sh,o["fill"],o["border"],o["lw"])
            elif k=="line":
                sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,pval(o["x1"]),pval(o["y1"]),pval(o["x2"]),pval(o["y2"]))
                style_shape(sh,None,o["color"],o["lw"],o["dash"])
            elif k=="text":
                for j,line in enumerate(o["lines"]):
                    yy=o["y"]+j*o["size"]*o["leading"]
                    tw=width(line,o["size"],o["bold"])
                    xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
                    sh=slide.shapes.add_textbox(pval(xx),pval(yy),pval(tw+8),pval(o["size"]*1.4))
                    tf=sh.text_frame; tf.clear(); tf.word_wrap=False
                    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                    p=tf.paragraphs[0]; p.space_before=Pt(0); p.space_after=Pt(0)
                    p.line_spacing=pval(o["size"]*o["leading"])
                    for f,txt in runs(line,o["bold"]):
                        run=p.add_run(); run.text=txt
                        run.font.name="Microsoft YaHei" if f=="CJK" else "Arial"
                        run.font.size=pval(o["size"]); run.font.bold=o["bold"]
                        run.font.color.rgb=ppt_color(o["color"])
                        if line.startswith("https://"):
                            run.hyperlink.address=line
                        props=run._r.get_or_add_rPr()
                        ea=OxmlElement("a:ea"); ea.set("typeface","Microsoft YaHei");props.append(ea)
        slide.notes_slide.notes_text_frame.text=scene.note+(("\n"+notes_source_label+scene.source) if scene.source else "")
    path=ROOT/filename
    prs.save(path)
    return path


def render_pdf(filename="LabOS_NLH_BVPE_System.pdf", title=None):
    path=ROOT/filename
    c=canvas.Canvas(str(path),pagesize=(W*SCALE,H*SCALE),pageCompression=1)
    c.setTitle(title or "LabOS：跨 PC 协同的非线性 Hall 与 BVPE 实验系统")
    c.setAuthor("LabOS")
    for scene in SLIDES:
        c.saveState(); c.scale(SCALE,SCALE)
        for o in scene.ops:
            k=o["kind"]
            if k in {"rect","circle","poly"}:
                if o["fill"]:c.setFillColorRGB(*rgb(o["fill"]))
                if o.get("border"):c.setStrokeColorRGB(*rgb(o["border"]))
                c.setLineWidth(o.get("lw",1)); c.setDash([5,4] if o.get("dash") else [])
                if k=="rect":
                    c.roundRect(o["x"],H-o["y"]-o["h"],o["w"],o["h"],o["radius"],
                                stroke=int(bool(o["border"])),fill=int(bool(o["fill"])))
                elif k=="circle":
                    c.circle(o["x"],H-o["y"],o["r"],stroke=int(bool(o["border"])),fill=int(bool(o["fill"])))
                else:
                    p=c.beginPath();p.moveTo(o["points"][0][0],H-o["points"][0][1])
                    for x,y in o["points"][1:]:p.lineTo(x,H-y)
                    p.close();c.drawPath(p,stroke=int(bool(o["border"])),fill=int(bool(o["fill"])))
            elif k=="line":
                c.setStrokeColorRGB(*rgb(o["color"])); c.setLineWidth(o["lw"])
                c.setDash([5,4] if o["dash"] else [])
                c.line(o["x1"],H-o["y1"],o["x2"],H-o["y2"])
            elif k=="text":
                c.setFillColorRGB(*rgb(o["color"]))
                for j,line in enumerate(o["lines"]):
                    yy=H-o["y"]-o["size"]*.93-j*o["size"]*o["leading"]
                    tw=width(line,o["size"],o["bold"])
                    xx=o["x"]+((o["w"]-tw)/2 if o["align"]=="center" else o["w"]-tw if o["align"]=="right" else 0)
                    if line.startswith("https://"):
                        c.linkURL(line,(xx,yy-2,xx+tw,yy+o["size"]),relative=1,thickness=0)
                    for f,txt in runs(line,o["bold"]):
                        c.setFont(f,o["size"]);c.drawString(xx,yy,txt)
                        xx+=pdfmetrics.stringWidth(txt,f,o["size"])
        c.restoreState();c.showPage()
    c.save()
    return path


def main():
    build_slides()
    for i,s in enumerate(SLIDES,1):s.footer(i,len(SLIDES))
    errors=[f"slide {i}: {e}" for i,s in enumerate(SLIDES,1) for e in s.errors]
    if errors:
        raise ValueError("\n".join(errors))
    pptx=render_pptx(); pdf=render_pdf()
    (ROOT/"slides_scene.json").write_text(json.dumps([dict(title=s.title,section=s.section,ops=s.ops,note=s.note,source=s.source)
                                                     for s in SLIDES],ensure_ascii=False,indent=2))
    script="# LabOS 跨 PC 协同实验系统讲解稿\n\n共 %d 页，建议讲解时间 15–20 分钟。\n\n"%len(SLIDES)
    for i,s in enumerate(SLIDES,1):
        script+=f"## {i:02d} · {s.title}\n\n{s.note}\n\n"
        if s.source:script+=f"参考：{s.source}\n\n"
    (ROOT/"讲解稿.md").write_text(script,encoding="utf-8")
    print(json.dumps(dict(slides=len(SLIDES),pptx=str(pptx),pdf=str(pdf),geometry_checks="PASS"),ensure_ascii=False))


if __name__=="__main__":
    main()
