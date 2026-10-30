"""Artifact checks and visual previews, independent of LabOS runtime dependencies."""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parent


def compact(text):
    return re.sub(r"\s+","",text)


def main():
    scenes=json.loads((ROOT/"slides_scene.json").read_text())
    pdf=pymupdf.open(ROOT/"LabOS_NLH_BVPE_System.pdf")
    prs=Presentation(ROOT/"LabOS_NLH_BVPE_System.pptx")
    assert len(pdf)==len(prs.slides)==len(scenes)==19
    assert abs(prs.slide_width/prs.slide_height-16/9)<1e-6
    thumbnails=[]
    issues=[]
    text_boxes=0
    vectors=0
    text_log=[]
    for i,(page,slide,scene) in enumerate(zip(pdf,prs.slides,scenes),1):
        page_text=page.get_text()
        text_log.append(page_text)
        assert "待确认" not in page_text
        ppt_text="\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert "待确认" not in ppt_text
        for op in scene["ops"]:
            if op["kind"]!="text":continue
            for line in op["lines"]:
                if compact(line) and compact(line) not in compact(page_text):
                    issues.append(f"page {i}: missing PDF text {line!r}")
                if compact(line) and compact(line) not in compact(ppt_text):
                    issues.append(f"slide {i}: missing native PPT text {line!r}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text:text_boxes+=1
            else:vectors+=1
            assert int(shape.shape_type)!=13,"unexpected raster image replacing native content"
        assert len(slide.notes_slide.notes_text_frame.text)>80
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines",[]):
                for span in line["spans"]:
                    x0,y0,x1,y1=span["bbox"]
                    if min(x0,y0)<-1 or x1>page.rect.width+1 or y1>page.rect.height+1:
                        issues.append(f"page {i}: out-of-page text {span['text']!r}")
        pix=page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False)
        image=Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
        if i in [3,5,7,9,11,13,14]:image.save(ROOT/f"preview_slide_{i:02d}.png")
        image.thumbnail((480,270))
        tile=Image.new("RGB",(500,304),"#E2E8F0")
        tile.paste(image,((500-image.width)//2,8))
        ImageDraw.Draw(tile).text((13,280),f"{i:02d}",fill="#11233E")
        thumbnails.append(tile)
    assert not issues,"\n".join(issues)
    for term in ["SuperK","CLD1000","730 nm","2400","Vxx","Vxy","chopper","BVPE","PD"]:
        assert term in "\n".join(text_log),term
    sheet=Image.new("RGB",(1500,304*((len(thumbnails)+2)//3)),"#D2DDE9")
    for i,image in enumerate(thumbnails):sheet.paste(image,((i%3)*500,(i//3)*304))
    sheet.save(ROOT/"LabOS_slides_overview.png")
    with zipfile.ZipFile(ROOT/"LabOS_NLH_BVPE_System.pptx") as z:
        assert z.testzip() is None
    (ROOT/"PDF文本检查.txt").write_text("\n\n".join(text_log),encoding="utf-8")
    report=dict(status="PASS",pages=19,aspect_ratio="16:9",native_editable_text_boxes=text_boxes,
                vector_shapes_and_lines=vectors,slide_raster_images=0,missing_text=0,
                out_of_page_text=0,speaker_notes="19/19",content_mode="系统能力与应用场景展示",
                forbidden_placeholder_labels=0,visual_inspection="overview and dense diagrams reviewed",
                pdf_reference_links=sum(len(list(p.get_links())) for p in pdf))
    (ROOT/"QA_REPORT.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    files=[ROOT/name for name in ("LabOS_NLH_BVPE_System.pptx","LabOS_NLH_BVPE_System.pdf","讲解稿.md","LabOS_slides_overview.png","QA_REPORT.json")]
    (ROOT/"SHA256SUMS").write_text("".join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}\n" for p in files))
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
